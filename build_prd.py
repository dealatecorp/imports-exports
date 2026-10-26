from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(27)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(12)
doc.core_properties.title = 'Import and Export Management Platform Product Requirements'
doc.core_properties.subject = 'Inventory warehouse loading unloading and shipment tracking'

def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t, 2)
def b(t): doc.add_paragraph(t, 'List Bullet')
def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)
def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=txt
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); pr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E5EAF0' if ri==0 else ('F6F7F9' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for run in para.runs:
                    run.font.size=Pt(10)
                    run.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.add_heading('Import and Export Management Platform',0)
doc.add_paragraph('Product Requirements Document', 'Subtitle')
p('Version 1.0  |  26 October 2026  |  Proposed scope for product review')
h('Product purpose')
p('Build a web application that gives an import and export organisation a shared view of warehouse stock, container loading, incoming cargo and order delivery. Managers control operations and approvals; employees record day-to-day warehouse and shipment activity. Two connected panels cover Export Operations for loading and dispatch, and Import Operations for arrival and unloading.')
h('The operating problem')
p('Warehouse quantities, container contents and shipment updates must remain consistent across teams. A shipment leaving the warehouse must reduce the correct stock once, and an arriving shipment must increase stock only after receipt. Users need to see what is available, what is committed, where cargo was last confirmed and what action is required next.')
h('Intended readers and decisions')
p('This document supports the organisation owner, operations managers, designers, developers and QA team. It defines the proposed first release, operational rules and acceptance criteria. The business owner should confirm the decisions on the final page before implementation estimates are committed.')
h('Product outcomes')
b('Maintain traceable inventory balances across warehouses and storage locations.')
b('Load and dispatch containers with verified item quantities and a complete manifest.')
b('Receive imported cargo with clear shortage, damage and overage handling.')
b('Show an ecommerce-style order timeline linked to shipment and container details.')
h('Scope convention')
p('All numbered requirements in this document are proposed MVP requirements unless explicitly marked as later scope. The phrase “items importing” within the loading panel is represented as “Loading and Dispatch.” Ocean checkpoints represent recorded events; they do not imply continuous GPS tracking.')

page('1 Users permissions and navigation')
h('Login and access')
p('AUTH 01 — Provide manager and employee login entry points backed by one identity system. The assigned account role determines permissions; selecting a login tab cannot grant access. Support password reset, logout, account deactivation and session expiry. Managers create or invite employees and assign warehouses and operational responsibilities.')
table(['Capability','Manager','Employee'],[
('View dashboards and records','All authorised operations','Assigned warehouses and shipments'),
('Products and warehouse setup','Create and edit','Read'),
('Receipts transfers and loading','Create and confirm','Record assigned operations'),
('Manual stock corrections','Approve or reject','Submit with reason'),
('Dispatch and exception closure','Approve and confirm','Prepare and submit'),
('Tracking and unloading events','Record and correct with audit','Record assigned events'),
('Users reports and audit history','Manage and view','Limited operational history')],[2.5,2.05,2.12])
p('AUTH 02 — Enforce permissions on every server request, including exports and attachments. Employees cannot read unrelated warehouse records by changing a URL or record identifier. Manager-created stock corrections require a reason and an audit entry; a second approver is a future policy option.')
h('Primary navigation')
b('Overview: operational summary, pending approvals and alerts.')
b('Export Operations: Ready to Load, Loading and Dispatch, Outbound Shipments.')
b('Import Operations: Incoming Orders, Shipment Details, Unloading and Receiving.')
b('Shared modules: Warehouse Inventory, Products, Warehouses, Reports and Settings.')
h('Dashboard requirements')
p('DASH 01 — Show warehouse on-hand, available, reserved and blocked quantities by SKU and unit; ready-to-load allocations; containers awaiting dispatch; active shipments; delayed shipments; arrivals due; and unresolved discrepancies. Do not combine incompatible units into one inventory quantity.')
p('DASH 02 — Every card opens the matching filtered records. Support warehouse, import/export direction, status and date filters. Display the selected timezone and last refresh time. Employees see only authorised data; managers see pending approvals and workload by warehouse.')

page('2 Warehouse inventory and stock integrity')
h('Stock definitions')
p('INV 01 — Track every SKU by warehouse, location and optional lot. On-hand includes available, reserved, staged, loaded-but-not-dispatched and blocked stock. Available equals on-hand minus reserved, staged, loaded-but-not-dispatched and blocked quantities. These buckets must be mutually exclusive. Dispatched cargo is tracked separately from warehouse on-hand.')
p('The Ready to Load screen is a filtered operational view of approved outbound allocations. It is not a second inventory balance. Imported stock becomes available only after receipt and acceptance; damaged or uninspected receipts remain blocked.')
h('Manual warehouse updates')
p('INV 02 — Allow employees to record receipts, put-away, location transfers and stock counts with SKU, quantity, unit, warehouse, location, date and notes. Each correction creates an adjustment request with old quantity, counted quantity, difference, reason and optional evidence. Apply the correction only after manager approval.')
p('INV 03 — Maintain an append-only movement ledger with movement type, signed quantity, source and destination, reference record, actor and timestamp. Correct posted mistakes using linked reversals or compensating movements. Do not overwrite or delete the original transaction.')
h('Inventory validation')
b('Reject zero or negative movement quantities and unsupported units. Store quantities at the configured SKU precision; do not mix pieces, kilograms and cartons without an explicit conversion.')
b('Prevent reservations and loading beyond available or allocated quantities. Use transactional updates so concurrent users cannot allocate the same stock twice.')
b('Repeated confirmation requests must create only one stock movement. Show a clear success or validation message, and preserve entered data after recoverable errors.')
b('Transfers within a warehouse change location balances without changing total stock. Transfers between warehouses use an in-transfer state until destination receipt.')
b('Keep products with transaction history archived rather than deleted. Require a SKU, product name and unit before any stock movement.')
h('Warehouse inventory screen')
p('INV 04 — Provide search by SKU and product name; filters by warehouse, location, lot and stock state; and columns for on-hand, available, reserved, staged, loaded and blocked. A product detail view shows movement history and linked orders. Managers can export the filtered list as CSV and configure low-stock thresholds by warehouse and SKU.')

page('3 Export loading and dispatch')
h('Outbound workflow')
p('Create outbound order → reserve stock → stage for loading → allocate to containers → confirm loading → manager dispatch → record transport milestones → confirm handover or delivery.')
p('EXP 01 — Create an outbound order with consignee, origin warehouse, destination, required date and SKU quantities. One order may be fulfilled across multiple shipments, and a shipment may contain multiple orders. Every loaded line must reference its order line and stock source.')
p('EXP 02 — Ready to Load lists approved quantities awaiting loading with SKU, warehouse, location, order, reserved quantity, loaded quantity and remaining quantity. Employees can stage and load partial quantities. Cancelling an unshipped allocation releases stock; loaded stock must be unloaded through a recorded reversal before release.')
h('Containers and cargo manifest')
p('EXP 03 — Create shipment and container records. Capture container number, type, seal number where applicable, carrier, vessel and voyage when known, departure and destination ports, planned departure and arrival, shipment type and assigned employees. Shipment type includes full-container, less-than-container or other configured cargo classification.')
p('EXP 04 — Allocate SKU quantities, lots, package counts, weight and volume to each container. Show planned and actual contents, container count and loading progress. Compare known weight and volume with configured limits and block dispatch when a configured limit is exceeded. Missing measurements must be visible; unsupported capacity checks must not appear as passed.')
h('Dispatch control')
p('EXP 05 — Before dispatch, validate the manifest, loaded quantities, required container identifiers, origin and destination, and the organisation’s configured document checklist. The manager confirms dispatch with an actual timestamp. Confirmation removes loaded quantities from warehouse on-hand exactly once and marks them as dispatched cargo.')
p('EXP 06 — Prevent ordinary edits to dispatched quantities. Route corrections through a manager-approved exception with linked stock movements when applicable. Cancelling a dispatched shipment must not automatically restock cargo that has not physically returned.')
h('Operational documents')
p('EXP 07 — Store attachments such as packing lists, invoices, bills of lading and loading photos against the shipment or container. Show document type, upload time and uploader. Generate a downloadable manifest of containers and cargo. Document requirements are configurable operational checks; customs filing and regulatory validation are outside MVP.')

page('4 Import unloading and order tracking')
h('Incoming workflow')
p('Register incoming order → link shipment and containers → track transport → record arrival → unload and inspect → receive accepted stock → resolve discrepancies → confirm delivery or close the order.')
p('IMP 01 — Register incoming orders with supplier, receiving party, destination warehouse, expected SKU quantities, expected arrival and reference numbers. Link each order line to one or more shipment/container allocations. Preserve expected, received, accepted, damaged and outstanding quantities separately.')
h('Shipment details')
p('IMP 02 — Display shipment reference, carrier, vessel and voyage when known, origin and destination ports, planned and actual dates, container count, container numbers and types, shipment type, commodity descriptions, package counts, weight, volume, seal details, attached documents and latest confirmed checkpoint. Container count is derived from linked container records.')
h('Unloading and receipt')
p('IMP 03 — Employees record unloading per container and item, with actual quantity, condition, storage location, timestamp and optional evidence. Support partial unloading and repeated receipt sessions. Receipt increases warehouse stock once; accepted stock becomes available and damaged or pending-inspection stock becomes blocked.')
p('IMP 04 — Compare actual receipt with expected quantities. Create a discrepancy for shortages, damage or overages. Managers approve overages and resolve discrepancies with an outcome and notes. Shortages remain outstanding until received or explicitly closed with an approved reason. Unknown products require a mapped SKU before inventory posting.')
h('Ecommerce style order timeline')
p('ORD 01 — Provide a simple stepper: Confirmed → Preparing → Shipped → In Transit → Arrived → Receiving → Completed. Show completed step timestamps, the current step, estimated next milestone, latest update and any delay. Show Out for Delivery and Delivered only when last-mile delivery is tracked.')
p('ORD 02 — An order detail page lists items, ordered and received quantities, linked shipments, containers and outstanding items. Split shipments have separate timelines and ETAs. Display Partially Received or Partially Delivered where appropriate; one completed container cannot mark the entire order complete.')
p('ORD 03 — MVP order tracking is available to authorised internal users. A customer portal or secure external tracking link is later scope. Visual familiarity with ecommerce tracking does not include shopping, checkout or payments.')

page('5 Shipment events and status rules')
h('Checkpoint tracking')
p('TRK 01 — Support configurable checkpoints such as origin warehouse dispatch, origin port, vessel departure, transshipment port, ocean waypoint, destination port, unloading and final handover. Each event contains shipment/container scope, event type, location, event time, recorded time, source, actor and notes. Store coordinates only when supplied.')
p('TRK 02 — MVP tracking uses manual updates from authorised employees and managers. Display the last confirmed location, source and update time. A route map is optional; if shown, label planned routes separately from confirmed positions. Never animate or infer live vessel positions from elapsed time.')
p('TRK 03 — Preserve the original ETA and record revised ETAs with reason and history. Highlight overdue milestones and stale updates using manager-configured thresholds. A delay or hold is an exception indicator and does not erase the underlying transport stage.')
table(['Record','Normal progression','Completion rule'],[
('Outbound cargo','Draft → Reserved → Staged → Loaded → Dispatched','Dispatch confirmed and stock posted once'),
('Transport','Planned → Departed → In Transit → Arrived','Arrival supported by a recorded event'),
('Inbound receipt','Expected → Unloading → Partially Received → Received','All expected quantities received or variance closed'),
('Order','Open → In Progress → Partial → Completed','All lines fulfilled or approved residual closure')],[1.15,3.0,2.52])
h('Corrections and exceptions')
p('TRK 04 — Validate transitions and retain a full event history. Late or out-of-order events may be added to history without automatically regressing the current status. Managers may correct or invalidate an event with a reason; keep both original and corrected records. Container-specific events affect only the linked cargo; shipment-wide updates must show their scope.')
p('TRK 05 — A shipment may have mixed container states. Summaries show counts by state and partial progress. Received, delivered and completed quantities are computed from line-level records rather than an editable summary label.')
h('Notifications and future integrations')
p('TRK 06 — Provide in-app alerts for low stock, pending approval, delays, arrival, discrepancies and stale tracking. Link alerts to records, deduplicate repeated alerts and allow acknowledgement. Email, SMS and messaging channels are later scope.')
p('Later scope — Integrate carrier or maritime tracking feeds after provider selection and coverage validation. Normalise provider events, deduplicate callbacks, retain source identifiers, retry failures and distinguish reported position from estimated position. Provider availability, cost and refresh frequency remain delivery dependencies.')

page('6 Data reporting and quality requirements')
h('Core data relationships')
p('DATA 01 — Use stable identifiers for users, roles, warehouses, locations, products, lots, orders, order lines, allocations, shipments, containers, cargo lines, stock movements, receipts, discrepancies, tracking events and attachments. An allocation connects an order line to cargo quantities, a shipment and a container. Prevent orphaned references and duplicate active container assignments.')
p('DATA 02 — Require explicit units for quantities, weight and volume. Store timestamps in UTC and display them in the organisation’s configured timezone with the timezone visible where needed. Use unique internal references even when external booking or bill-of-lading numbers are reused.')
h('Reports')
p('REP 01 — Managers can filter and export inventory balances, stock movements, ready-to-load quantities, container manifests, active and delayed shipments, incoming orders, receipt discrepancies and order fulfilment. Exports must match the applied filters and permissions. Include generation time, timezone and units.')
h('Security and reliability')
b('NFR 01 — Enforce server-side role and warehouse access, encrypted transport, secure password storage or a managed identity provider, session expiry and rate limits on login. Require manager MFA for production launch.')
b('NFR 02 — Audit sign-in relevant security events, role changes, approvals, stock postings, dispatch, receipts and tracking corrections. Record actor, time, action and affected entity; audit records are not editable through the application.')
b('NFR 03 — Validate file size and type, scan uploads before access and use permission-controlled downloads. Avoid exposing documents through permanent public URLs.')
b('NFR 04 — Use atomic stock transactions, duplicate-request protection and conflict messages. Refresh changed records before allowing a stale edit to overwrite operational data.')
b('NFR 05 — Proposed service targets: 99.5% monthly availability, recovery point within 24 hours and recovery time within 4 hours. Schedule backups and verify restoration before launch; targets require infrastructure approval.')
h('Performance and usability')
p('NFR 06 — Proposed load-test baseline: 50 concurrent users, 10 warehouses, 10,000 SKUs and 100,000 stock movements. At this baseline, 95% of common list/detail requests should complete within 2 seconds and stock confirmations within 3 seconds, excluding file transfer and external providers. Use pagination and indexed search.')
p('NFR 07 — Support desktop and tablet warehouse workflows and mobile order tracking. Provide keyboard operation, visible focus, labelled fields, readable contrast and status text independent of colour. Include loading, empty, error and retry states. MVP requires an internet connection; offline posting is later scope.')

page('7 Acceptance criteria and release validation')
p('The following scenarios form the minimum business acceptance suite. QA must also test permission failures, invalid inputs, partial completion, retries and concurrent updates across both panels.')
table(['ID','Scenario and pass condition'],[
('AC 01','An employee requests another warehouse record or export by identifier. Access is denied and no restricted content is returned.'),
('AC 02','With 100 units on-hand and 20 blocked, reserve 60 units. Available becomes 20. A second request to reserve 21 units fails without changing balances.'),
('AC 03','Load 40 of the 60 reserved units. Reserved becomes 20, loaded becomes 40, on-hand remains 100 and available remains 20. Dispatch the 40 units: on-hand becomes 60 and loaded becomes zero.'),
('AC 04','Retry the same dispatch or receipt confirmation after a timeout. Exactly one stock movement and one business confirmation exist.'),
('AC 05','An employee submits a stock correction. Balances remain unchanged until approval; approval records the difference, actor, reason and audit history.'),
('AC 06','For 100 expected units, receive 90 accepted and 5 damaged. Post 95 on-hand, including 5 blocked, and keep a shortage of 5 outstanding with an open discrepancy.'),
('AC 07','Split an order across two containers. Receive one container; show partial progress, its actual receipt and the other container’s ETA. Do not complete the order.'),
('AC 08','Add an older departure event after an arrival event. Both appear in chronological history; the current status remains Arrived.'),
('AC 09','Cancel a reserved but unshipped allocation. Release its reservation. Cancellation after dispatch does not create warehouse stock.'),
('AC 10','Two employees reserve the last available units simultaneously. Only valid quantities commit; total reservations never exceed eligible stock.'),
('AC 11','Import and export dashboards reconcile with filtered detail pages and the movement ledger. Quantities remain grouped by compatible units.'),
('AC 12','Complete an order with an unresolved shortage. Completion is blocked until receipt or manager-approved residual closure is recorded.')],[.65,6.02])
h('Launch gate')
p('All acceptance scenarios pass; no unresolved critical stock-integrity or access-control defects remain. Complete a restore drill, approved opening-stock reconciliation, manager and employee training, and end-to-end testing from outbound reservation to dispatch and from incoming order to receipt. Operations signs off the configured milestones and document checklist.')

page('8 Delivery scope and business decisions')
h('Release sequence')
b('Foundation: confirm workflows and data definitions; design navigation and role permissions; create product, warehouse and user records; validate opening-stock import or entry.')
b('Inventory and export: deliver movement ledger, manual updates and approvals, ready-to-load allocations, container loading, manifests and dispatch.')
b('Import and visibility: deliver incoming orders, unloading and receipts, discrepancy handling, manual checkpoints, timelines, dashboards and reports.')
b('Pilot and launch: test with one operational team and representative shipments, reconcile inventory, resolve launch defects and roll out to the remaining warehouses.')
p('Each stage must satisfy its applicable acceptance criteria before the next operational rollout. Calendar estimates require team capacity, integration decisions and validated data volumes.')
h('Later scope')
p('Carrier and AIS integrations; customer self-service tracking; barcode or QR scanning; offline warehouse work; purchase and sales accounting; invoicing and payments; customs filing; IoT sensors; advanced route optimisation; mobile apps; and multi-organisation SaaS tenancy. These additions need separate requirements and estimates.')
h('Decisions required before implementation')
table(['Decision','Proposed starting position'],[
('Organisation and access','One organisation with multiple warehouses; managers define employee assignments.'),
('Tracking audience','Internal manager and employee tracking first; confirm whether customers need launch access.'),
('Ocean checkpoint source','Manual events for MVP; select provider and budget if automatic tracking is required at launch.'),
('Stock ownership','Track operational warehouse custody; confirm consignment and customer-owned stock requirements.'),
('Receipt and delivery completion','Warehouse receipt completes inbound fulfilment unless last-mile delivery is explicitly in scope.'),
('Product complexity','Confirm lot, expiry, serial number, hazardous cargo and unit-conversion requirements.'),
('Operating policy','Confirm approval rules, mandatory documents, timezone, retention period and delay thresholds.'),
('Migration and scale','Confirm warehouse count, user count, opening balances, existing spreadsheets and import format.')],[2.0,4.67])
h('Success measures')
p('During the pilot, measure physical-count variance by SKU, time from receipt to inventory posting, percentage of shipments with current tracking, unresolved discrepancy age and time to find an order’s current status. Establish the baseline in the first two weeks, then agree improvement targets. Stock-changing actions must always retain an actor and traceable reference.')

out=ROOT/'Import and Export Management PRD.docx'
doc.save(out)
print(out)
